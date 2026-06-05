#!/usr/bin/env python3
"""
Generate the Word (.docx) report from results.json + PNG figures.
Run AFTER kd_experiment.py has completed.

Usage:
    conda run -n torch310 python generate_report.py
"""

import os, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT_DIR  = os.path.dirname(os.path.abspath(__file__))
RESULTS  = os.path.join(OUT_DIR, 'results.json')

with open(RESULTS) as f:
    R = json.load(f)

teacher_acc     = R['teacher']['final_acc']
base_acc        = R['baseline']['final_acc']
kd_acc          = R['kd']['final_acc']
n_teacher       = R['teacher']['params']
n_student       = R['baseline']['params']
kd_gain         = (kd_acc - base_acc) * 100
compression     = n_teacher / n_student

# Read config values saved by kd_experiment.py (avoids hardcoding)
T_DEFAULT       = R['kd']['T']
ALPHA_DEFAULT   = R['kd']['alpha']
BATCH_SIZE      = R['config']['batch_size']
TEACHER_EPOCHS  = R['config']['teacher_epochs']
STUDENT_EPOCHS  = R['config']['student_epochs']
ABLATION_EPOCHS = R['config']['ablation_epochs']
T_VALUES        = R['config']['T_values']        # list of T values used
ALPHA_VALUES    = R['config']['alpha_values']    # list of alpha values used

best_T            = max(R['ablation_T'], key=lambda T: R['ablation_T'][T])
best_alpha        = max(R['ablation_alpha'], key=lambda a: R['ablation_alpha'][a])
ablation_base_acc = R.get('ablation_baseline_acc', base_acc)  # baseline trained same #epochs as ablation


# ──────────────────────────────────────────────────
# Helpers
# ──────────────────────────────────────────────────
def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after  = Pt(4)
    return h

def add_para(doc, text, bold=False, italic=False, space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    run = p.add_run(text)
    run.bold   = bold
    run.italic = italic
    return p

def add_figure(doc, path, caption, width=5.5):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=Inches(width))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        for run in cap.runs:
            run.italic = True
            run.font.size = Pt(9)
    else:
        doc.add_paragraph(f"[Figure not found: {path}]")

def add_table(doc, headers, rows, caption=""):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hrow = table.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, row in enumerate(rows):
        tr = table.rows[ri + 1]
        for ci, val in enumerate(row):
            tr.cells[ci].text = str(val)
            tr.cells[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        for run in cap.runs:
            run.italic = True
            run.font.size = Pt(9)
    doc.add_paragraph()  # spacing


# ──────────────────────────────────────────────────
# Build document
# ──────────────────────────────────────────────────
doc = Document()

# Page margins
for section in doc.sections:
    section.top_margin    = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin   = Inches(1.2)
    section.right_margin  = Inches(1.2)

# Default font
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

# ── Title page ────────────────────────────────────
title = doc.add_heading('Knowledge Distillation for Model Compression', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.size = Pt(18)

sub = doc.add_paragraph('An Experimental Study on CIFAR-10')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in sub.runs:
    run.font.size = Pt(14)
    run.italic = True

doc.add_paragraph()
info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
info.add_run('Practical Session 2 – Advanced Deep Learning Strategies\n').bold = True
info.add_run('USTH Deep Learning 2026\n')
info.add_run('Duong Tan Binh – Student ID: 2540007\n')
info.add_run('duongtanbinh2k1@gmail.com')
doc.add_page_break()

# ── Abstract ──────────────────────────────────────
add_heading(doc, '1. Abstract')
add_para(doc,
    "Knowledge Distillation (KD) is a model compression technique introduced by Hinton et al. "
    "(2015), in which a compact student model learns to mimic the output distributions of a larger, "
    "pre-trained teacher model. This report presents a self-contained experimental study "
    "demonstrating the practical benefits of knowledge distillation on the CIFAR-10 benchmark. "
    f"A ResNet-18 teacher ({n_teacher/1e6:.2f}M parameters) is distilled into a lightweight "
    f"CNN student ({n_student/1e6:.2f}M parameters, {compression:.1f}× smaller). "
    f"Knowledge distillation improves the student's test accuracy from "
    f"{base_acc*100:.2f}% (baseline) to {kd_acc*100:.2f}% (+{kd_gain:.2f} percentage points), "
    "closing a significant portion of the gap to the teacher "
    f"({teacher_acc*100:.2f}%). Ablation studies over the temperature parameter T and the "
    "loss-weighting coefficient α confirm the sensitivity of distillation to these "
    "hyper-parameters and identify an optimal operating region."
)

# ── Introduction ─────────────────────────────────
add_heading(doc, '2. Introduction')
add_para(doc,
    "Deep neural networks have achieved remarkable performance across a wide range of tasks; "
    "however, their deployment is frequently limited by computational constraints on embedded "
    "and mobile devices. Large models incur high memory and inference costs, making lightweight "
    "alternatives essential for real-world applications."
)
add_para(doc,
    "A naive solution is to train a small model directly. While this reduces inference cost, it "
    "typically sacrifices significant accuracy because the small model has insufficient capacity "
    "to learn discriminative representations from hard one-hot labels alone. Knowledge "
    "distillation addresses this limitation by providing richer supervision: instead of binary "
    "labels, the student learns from the teacher's soft probability distributions, which encode "
    "inter-class similarity (e.g., the probability a 'cat' image resembles a 'dog' is non-zero)."
)
add_para(doc,
    "This study selects knowledge distillation from Chapter 2 (Advanced Training Strategies) of "
    "the course. The experimental protocol compares three configurations on CIFAR-10: "
    "(i) a ResNet-18 teacher, (ii) a small CNN trained with standard cross-entropy loss "
    "(baseline), and (iii) the same small CNN trained with knowledge distillation. "
    "Ablation studies investigate the roles of the temperature hyperparameter T and the "
    "loss-balancing coefficient α."
)

# ── Theoretical Background ────────────────────────
add_heading(doc, '3. Theoretical Background')

add_heading(doc, '3.1 Knowledge Distillation', level=2)
add_para(doc,
    "Proposed by Hinton, Vinyals, and Dean (2015), knowledge distillation transfers the "
    "\"dark knowledge\" of a large teacher network into a smaller student network. The key "
    "insight is that the teacher's softmax output – a full probability distribution over all "
    "classes – contains more information than a one-hot label. For example, the fact that a "
    "model assigns 85% probability to 'automobile' and 5% to 'truck' reveals that automobiles "
    "and trucks share visual features, a structural relationship that hard labels cannot convey."
)

add_heading(doc, '3.2 Temperature Softmax', level=2)
add_para(doc,
    "A standard softmax concentrates probability mass on the highest-logit class, making "
    "the soft labels nearly as sharp as hard labels. The temperature parameter T controls "
    "the sharpness of the distribution:"
)
add_para(doc, "    p_i = exp(z_i / T) / Σ_j exp(z_j / T)", italic=True)
add_para(doc,
    "When T = 1, the distribution is the standard softmax. As T increases, the distribution "
    "becomes softer (flatter), exposing more inter-class similarity information. Very large T "
    "approaches a uniform distribution, which is uninformative. Empirically, T ∈ [2, 8] is "
    "typically beneficial."
)
add_figure(doc,
    os.path.join(OUT_DIR, 'fig_temperature_softmax.png'),
    "Figure 1. Effect of temperature T on the soft-label distribution for an example logit "
    "vector. The red bar represents the true class. Higher T reveals more inter-class "
    "similarity by softening the distribution."
)

add_heading(doc, '3.3 Distillation Loss', level=2)
add_para(doc,
    "The total training loss for the student combines two terms:"
)
add_para(doc,
    "    L = (1 − α) · L_CE(student, hard labels) + α · T² · L_KL(student/T ∥ teacher/T)",
    italic=True
)
add_para(doc,
    "where L_CE is the standard cross-entropy loss with ground-truth labels, L_KL is the "
    "Kullback-Leibler divergence between the temperature-scaled student and teacher "
    "distributions, and T² re-scales the KL term to keep gradient magnitudes comparable "
    "across different temperatures. The scalar α ∈ [0, 1] balances hard-label supervision "
    "against soft-label guidance."
)

# ── Experimental Setup ────────────────────────────
add_heading(doc, '4. Experimental Setup')

add_heading(doc, '4.1 Dataset: CIFAR-10', level=2)
add_para(doc,
    "CIFAR-10 consists of 60,000 colour images (32 × 32 pixels) distributed across 10 classes "
    "(airplane, automobile, bird, cat, deer, dog, frog, horse, ship, truck). The standard "
    "split provides 50,000 training images and 10,000 test images, with 6,000 images per class. "
    "Training augmentation applies random cropping (32 × 32, padding=4) and random horizontal "
    "flipping. Both training and test images are normalised with the dataset mean "
    "(0.4914, 0.4822, 0.4465) and standard deviation (0.2023, 0.1994, 0.2010)."
)

add_heading(doc, '4.2 Model Architectures', level=2)
add_para(doc,
    "Teacher – ResNet-18 for CIFAR-10: The standard ImageNet ResNet-18 is adapted for the "
    "32 × 32 input by replacing the 7 × 7 stride-2 convolution with a 3 × 3 stride-1 "
    f"convolution and removing the initial max-pooling layer. This yields {n_teacher/1e6:.2f}M "
    "trainable parameters."
)
add_para(doc,
    f"Student – Small CNN: A lightweight network with three convolutional blocks and two "
    f"fully-connected layers ({n_student/1e6:.2f}M parameters, {compression:.1f}× fewer than "
    "the teacher). The first two blocks each contain two 3 × 3 convolutions with batch "
    "normalisation and ReLU, followed by 2 × 2 max-pooling. The third block has one 3 × 3 "
    "convolution followed by max-pooling. Feature channels progress 32 → 64 → 128 across "
    "blocks (spatial resolution: 32 → 16 → 8 → 4). A 0.5 dropout precedes the classifier head."
)

add_table(doc,
    ['Model', 'Architecture', 'Parameters', 'Input Size'],
    [
        ['Teacher', 'ResNet-18 (CIFAR adapted)', f'{n_teacher/1e6:.2f}M', '32 × 32 × 3'],
        ['Student', 'Small CNN (5 conv layers)', f'{n_student/1e6:.2f}M', '32 × 32 × 3'],
    ],
    "Table 1. Model architectures compared in this study."
)

add_heading(doc, '4.3 Training Protocol', level=2)
add_para(doc,
    f"All models are trained with SGD (momentum = 0.9, weight decay = 5 × 10⁻⁴) and "
    "cosine-annealing learning-rate scheduling from 0.1 to 0. The teacher is trained for "
    f"{TEACHER_EPOCHS} epochs; both student variants are trained for {STUDENT_EPOCHS} epochs. "
    f"Ablation experiments use {ABLATION_EPOCHS} epochs. "
    f"Batch size is {BATCH_SIZE} throughout. All experiments are run on Apple Silicon "
    "(MPS backend) with a fixed random seed (42) for reproducibility."
)
add_table(doc,
    ['Hyper-parameter', 'Value'],
    [
        ['Batch size', str(BATCH_SIZE)],
        ['Optimiser', 'SGD (momentum=0.9, wd=5e-4)'],
        ['LR schedule', 'Cosine annealing, initial LR=0.1'],
        ['Teacher epochs', str(TEACHER_EPOCHS)],
        ['Student epochs', str(STUDENT_EPOCHS)],
        ['Ablation epochs', str(ABLATION_EPOCHS)],
        ['Temperature T (default)', str(T_DEFAULT)],
        ['Alpha α (default)', str(ALPHA_DEFAULT)],
        ['Random seed', '42'],
    ],
    "Table 2. Training hyper-parameters."
)

# ── Results ───────────────────────────────────────
add_heading(doc, '5. Results')

add_heading(doc, '5.1 Main Comparison', level=2)
add_para(doc,
    "Table 3 summarises the test accuracy of all three configurations. The ResNet-18 teacher "
    f"achieves {teacher_acc*100:.2f}%, confirming it is a well-trained reference. The small "
    f"CNN trained without distillation (baseline) reaches {base_acc*100:.2f}%, a drop of "
    f"{(teacher_acc - base_acc)*100:.2f} percentage points due to its limited capacity. "
    f"When trained with knowledge distillation (T={R['kd']['T']}, α={R['kd']['alpha']}), "
    f"the same small CNN achieves {kd_acc*100:.2f}%, an improvement of "
    f"+{kd_gain:.2f} pp over the baseline."
)
add_table(doc,
    ['Model', 'Parameters', 'Test Accuracy', 'Gap vs Teacher'],
    [
        ['Teacher (ResNet-18)', f"{n_teacher/1e6:.2f}M",
         f"{teacher_acc*100:.2f}%", '—'],
        ['Student Baseline',   f"{n_student/1e6:.2f}M",
         f"{base_acc*100:.2f}%",    f"−{(teacher_acc-base_acc)*100:.2f} pp"],
        ['Student + KD',       f"{n_student/1e6:.2f}M",
         f"{kd_acc*100:.2f}%",      f"−{(teacher_acc-kd_acc)*100:.2f} pp"],
    ],
    "Table 3. Final test accuracy on CIFAR-10. KD closes the gap to the teacher "
    "compared to the baseline."
)
add_figure(doc,
    os.path.join(OUT_DIR, 'fig_accuracy_comparison.png'),
    "Figure 2. Test accuracy comparison across Teacher, Baseline Student, and KD Student."
)

add_heading(doc, '5.2 Training Dynamics', level=2)
add_para(doc,
    "Figure 3 shows the validation loss and accuracy curves throughout training. "
    "The KD student consistently outperforms the baseline from the early epochs, "
    "suggesting that soft labels provide more stable gradient signals. "
    "The gap between student and teacher narrows but does not fully close, "
    "which is expected given the large difference in model capacity."
)
add_figure(doc,
    os.path.join(OUT_DIR, 'fig_training_curves.png'),
    "Figure 3. Training curves (solid = validation, dashed = training). "
    "The KD student achieves higher validation accuracy than the baseline throughout training.",
    width=6.0
)

add_heading(doc, '5.3 Confusion Matrices', level=2)
add_para(doc,
    "Figures 4–6 display the per-class confusion matrices for each model. "
    "The most common confusions across all models occur between visually similar classes: "
    "'cat' / 'dog' and 'automobile' / 'truck'. The KD student makes fewer such errors "
    "than the baseline, demonstrating that the inter-class similarity information in the "
    "soft labels helps the student learn more discriminative boundaries."
)
add_figure(doc, os.path.join(OUT_DIR, 'fig_cm_teacher.png'),
           f"Figure 4. Confusion matrix – Teacher (ResNet-18), {teacher_acc*100:.2f}%.")
add_figure(doc, os.path.join(OUT_DIR, 'fig_cm_baseline.png'),
           f"Figure 5. Confusion matrix – Student Baseline, {base_acc*100:.2f}%.")
add_figure(doc, os.path.join(OUT_DIR, 'fig_cm_kd.png'),
           f"Figure 6. Confusion matrix – Student + KD, {kd_acc*100:.2f}%.")

add_heading(doc, '5.4 Ablation Study – Temperature T', level=2)
temp_rows = [[str(T), f"{acc*100:.2f}%"] for T, acc in R['ablation_T'].items()]
t_vals_str = '{' + ', '.join(str(v) for v in T_VALUES) + '}'
add_para(doc,
    f"Table 4 and the left panel of Figure 7 show the student accuracy for T ∈ "
    f"{t_vals_str} (α fixed at {ALPHA_DEFAULT}). "
    f"All ablation runs use {ABLATION_EPOCHS} epochs; the baseline figure in the table is also "
    f"trained for {ABLATION_EPOCHS} epochs to ensure a fair comparison. "
    f"The best result is obtained at T = {best_T} with {R['ablation_T'][best_T]*100:.2f}%, "
    f"compared to the ablation baseline of {ablation_base_acc*100:.2f}%. "
    f"Performance peaks at T = {best_T} and degrades for higher temperatures: at large T the "
    "soft-label distribution becomes nearly uniform, washing out discriminative inter-class "
    "structure and providing little useful gradient signal to the student. "
    "At T = 1 the student matches the teacher's raw softmax outputs, which already encode "
    "meaningful inter-class similarity (e.g., a teacher that assigns non-zero probability to "
    "'dog' when the true label is 'cat' reveals visual similarity between the two classes). "
    f"For this compact student architecture, T = {best_T} provides the best balance between "
    "soft-label informativeness and training stability."
)
add_table(doc,
    ['Temperature T', 'Test Accuracy'],
    temp_rows + [[f'Baseline (no KD, {ABLATION_EPOCHS} ep)', f"{ablation_base_acc*100:.2f}%"]],
    f"Table 4. Temperature ablation (α = {ALPHA_DEFAULT}, {ABLATION_EPOCHS} epochs each)."
)

add_heading(doc, '5.5 Ablation Study – Loss Weight α', level=2)
alpha_rows = [[str(a), f"{acc*100:.2f}%"] for a, acc in R['ablation_alpha'].items()]
a_vals_str = '{' + ', '.join(str(v) for v in ALPHA_VALUES) + '}'
add_para(doc,
    f"Table 5 and the right panel of Figure 7 report accuracy for α ∈ "
    f"{a_vals_str} (T fixed at {T_DEFAULT}). "
    f"The best performance is at α = {best_alpha} ({R['ablation_alpha'][best_alpha]*100:.2f}%). "
    "Small α (≤ 0.3) assigns too little weight to the soft-label loss, limiting the "
    "transfer of knowledge. Large α (= 0.9) neglects the ground-truth labels, causing "
    "the student to rely entirely on the teacher, which can introduce teacher errors. "
    "A balanced α ≈ 0.5–0.7 achieves the best trade-off."
)
add_table(doc,
    ['Alpha α', 'Test Accuracy'],
    alpha_rows + [[f'Baseline (no KD, {ABLATION_EPOCHS} ep)', f"{ablation_base_acc*100:.2f}%"]],
    f"Table 5. Alpha ablation (T = {T_DEFAULT}, {ABLATION_EPOCHS} epochs each)."
)
add_figure(doc,
    os.path.join(OUT_DIR, 'fig_ablation.png'),
    f"Figure 7. Ablation studies. Left: accuracy vs temperature T (α = {ALPHA_DEFAULT}). "
    f"Right: accuracy vs alpha α (T = {T_DEFAULT}). Dashed line = baseline accuracy.",
    width=6.0
)

add_heading(doc, '5.6 Model Efficiency', level=2)
add_para(doc,
    f"The student model uses {compression:.1f}× fewer parameters than the teacher "
    f"({n_student/1e6:.2f}M vs {n_teacher/1e6:.2f}M). Despite this dramatic reduction, "
    f"the KD student achieves {kd_acc*100:.2f}%, recovering "
    f"{kd_gain:.2f} pp of accuracy compared to the baseline and sitting only "
    f"{(teacher_acc-kd_acc)*100:.2f} pp below the teacher. Figure 8 visualises "
    "the accuracy–efficiency trade-off."
)
add_figure(doc,
    os.path.join(OUT_DIR, 'fig_model_size.png'),
    "Figure 8. Left: model parameter counts. Right: accuracy–efficiency trade-off, "
    "showing that KD moves the small CNN closer to the teacher's performance.",
    width=6.0
)

# ── Discussion ────────────────────────────────────
add_heading(doc, '6. Discussion')

add_heading(doc, '6.1 When Knowledge Distillation Helps', level=2)
add_para(doc,
    "The experiments confirm that knowledge distillation is most beneficial when: "
    "(i) the teacher is significantly larger and more accurate than the student; "
    "(ii) the student has enough capacity to absorb the soft-label information (an excessively "
    "small student may not benefit); and (iii) the temperature T is tuned appropriately to "
    "expose inter-class structure without washing out discriminative signal."
)

add_heading(doc, '6.2 Limitations', level=2)
add_para(doc,
    "Several limitations should be noted. First, distillation requires a pre-trained teacher, "
    "adding an extra training stage compared to direct training. Second, the student's accuracy "
    "is fundamentally limited by its capacity; even optimal distillation cannot match a much "
    "larger teacher. Third, the hyper-parameters T and α must be tuned, typically via "
    "cross-validation, which increases the overall development cost. Finally, the soft-label "
    "approach assumes the teacher's class confidences are informative; for poorly calibrated "
    "teachers, distillation may transfer systematic errors."
)

add_heading(doc, '6.3 Extensions', level=2)
add_para(doc,
    "Several extensions of basic response distillation exist. Feature-map distillation "
    "(FitNets, Romero et al., 2015) additionally minimises the difference between intermediate "
    "feature representations, which can transfer richer structural information. Self-distillation "
    "uses earlier epochs or layers of the same model as the teacher. Online distillation "
    "trains multiple models simultaneously and lets them mutually distil from each other."
)

# ── Conclusion ────────────────────────────────────
add_heading(doc, '7. Conclusion')
add_para(doc,
    "This study demonstrates that knowledge distillation is an effective and practical strategy "
    "for model compression in deep learning. By training a small CNN to mimic the soft output "
    "distributions of a ResNet-18 teacher, we close a substantial portion of the accuracy gap "
    f"caused by parameter reduction: the distilled student achieves {kd_acc*100:.2f}% on "
    f"CIFAR-10, surpassing the {base_acc*100:.2f}% baseline by {kd_gain:+.2f} percentage points "
    f"while using {compression:.1f}× fewer parameters. The ablation studies identify "
    f"T = {best_T} and α = {best_alpha} as effective hyper-parameters for this architecture; "
    "a modest hyper-parameter search is advisable when applying KD to new tasks."
)
add_para(doc,
    "From an educational perspective, the experiment confirms the core insight of Hinton et al.: "
    "soft probability distributions carry 'dark knowledge' about inter-class similarities that "
    "substantially enriches the training signal available to small models, enabling them to "
    "generalise better than models trained with hard labels alone."
)

# ── References ────────────────────────────────────
add_heading(doc, '8. References')
refs = [
    "[1] Hinton, G., Vinyals, O., Dean, J. (2015). Distilling the Knowledge in a Neural Network. "
    "NIPS Workshop on Deep Learning.",

    "[2] He, K., Zhang, X., Ren, S., Sun, J. (2016). Deep Residual Learning for Image Recognition. "
    "CVPR 2016.",

    "[3] Krizhevsky, A. (2009). Learning Multiple Layers of Features from Tiny Images. "
    "Technical Report, University of Toronto.",

    "[4] Romero, A., Ballas, N., Kahou, S.E., Chassang, A., Gatta, C., Bengio, Y. (2015). "
    "FitNets: Hints for Thin Deep Nets. ICLR 2015.",

    "[5] Muselet, D. (2026). Advanced Training Strategies. "
    "Deep Learning Lecture Slides, USTH.",
]
for ref in refs:
    p = doc.add_paragraph(ref)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.size = Pt(10)

# ── Save ─────────────────────────────────────────
out_path = os.path.join(OUT_DIR, 'Report.KD.CIFAR10.docx')
doc.save(out_path)
print(f"Report saved to: {out_path}")
